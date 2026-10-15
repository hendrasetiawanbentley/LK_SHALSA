"""Rasio kertas kerja: rumus aman, perhitungan Y/Y-1 dari PDF, dan definisi template repo."""
from pathlib import Path

import pytest
from openpyxl import load_workbook

from sikuli import rasio

ROOT = Path(__file__).resolve().parents[1]


def angka(entity, kode, berjalan, lalu=None, skala=1.0):
    baris = [{"entity_key": entity, "kode_akun": kode, "periode": "berjalan", "nilai": berjalan, "skala": skala}]
    if lalu is not None:
        baris.append({"entity_key": entity, "kode_akun": kode, "periode": "lalu", "nilai": lalu, "skala": skala})
    return baris


def test_evaluasi_dan_akun_hilang():
    nilai = {"a": 10.0, "b": 4.0, "c": None}.get
    assert rasio.evaluasi("(a - opsional(c)) / rata2(a, b)", nilai) == pytest.approx(10 / 7)
    with pytest.raises(rasio.AkunHilang):
        rasio.evaluasi("a / c", nilai)


def test_rumus_berbahaya_ditolak():
    with pytest.raises(ValueError):
        rasio._cek_rumus("__import__('os').system('x')")
    with pytest.raises(ValueError):
        rasio._cek_rumus("a.b")


def test_hitung_rasio_dari_pdf_saja():
    definisi = [
        {"no": 1, "entitas": "MI", "satuan": "persen", "rumus": "laba / aset", "teks": "Laba / Aset x 100%",
         "kategori": "Profitabilitas", "field_key": "ROA", "dtype": "number", "description": ""},
        {"no": 21, "entitas": "MI", "satuan": "nominal", "rumus": "modal", "teks": "Modal >= 25 M", "ambang": 25e9,
         "kategori": "Permodalan", "field_key": "MIKU 1", "dtype": "number", "description": ""},
        {"no": 9, "entitas": "MI", "satuan": "kali", "rumus": "laba / beban", "teks": "x",
         "kategori": "K", "field_key": "ICR", "dtype": "number", "description": ""},
        {"no": 99, "entitas": "MI", "satuan": "persen", "rumus": None, "teks": "",
         "kategori": "K", "field_key": "Baru", "dtype": "number", "description": ""},
    ]
    pdf = angka("E", "laba", 5, 4, skala=1e6) + angka("E", "aset", 100, 80, skala=1e6) + angka("E", "modal", 26e9)
    hasil = {r["field_key"]: r for r in rasio.hitung_rasio(definisi, pdf, {"E": "MI"}, {"laba": "Laba", "aset": "Total Aset"})}
    roa = hasil["ROA"]
    assert roa["rasio_y"] == pytest.approx(5.0) and roa["rasio_y1"] == pytest.approx(5.0)
    assert roa["nilai_y"] == "Laba = 5.000.000,00\nTotal Aset = 100.000.000,00"
    assert roa["nilai_y1"] == "Laba = 4.000.000,00\nTotal Aset = 80.000.000,00"
    assert roa["hasil_rasio"] == "Y-1: 5,00%\nY: 5,00%"
    assert hasil["MIKU 1"]["hasil_rasio"].endswith("Y: 26.000.000.000 (memenuhi)")
    assert hasil["ICR"]["rasio_y"] is None and "akun beban tidak tersedia di PDF" in hasil["ICR"]["keterangan"]
    assert hasil["Baru"]["keterangan"] == "rumus belum didefinisikan"


def test_definisi_template_repo_dan_kk(tmp_path):
    definisi = rasio.muat_rasio(ROOT / "config" / "rasio_kk_template.xlsx", ROOT / "config" / "rasio.yaml")
    kunci = {(d["no"], d["entitas"]) for d in definisi}
    assert {(1, "EPP"), (20, "EPP"), (14, "MI"), (30, "RD")} <= kunci
    assert not kunci & {(21, "MI"), (22, "MI"), (25, "RD")}          # bertanda "tidak dipakai" di template
    assert all(d["rumus"] for d in definisi if d["entitas"] != "DES")
    baris = rasio.hitung_rasio(definisi, angka("X_A_1", "total_aset", 10.0), {"X_A_1": "EPP", "Y_B_2": "MI"})
    rasio.tulis_kk(baris, tmp_path / "kk.xlsx")
    wb = load_workbook(tmp_path / "kk.xlsx")
    assert wb.sheetnames == ["X_A", "Y_B", "Semua"]
    assert [c.value for c in wb["X_A"][1]][:10] == rasio.KOLOM_KK


def test_nama_sheet_unik():
    assert rasio.nama_sheet(["RD_STA01_Berbasis_X", "RD_STA01_STAR_Y", "MI_CC002_LK"]) == {
        "RD_STA01_Berbasis_X": "RD_STA01_Berbasis", "RD_STA01_STAR_Y": "RD_STA01_STAR", "MI_CC002_LK": "MI_CC002"}


def test_kelompok_periode_dan_batas_maks_untuk_kriteria_tambahan():
    definisi = [{"no": 31, "entitas": "DES", "satuan": "persen", "rumus": "utang_bunga / total_aset", "teks": "RU",
                 "kategori": "DES", "field_key": "RU Periode 1", "dtype": "number", "description": "",
                 "ambang_maks": 45, "periode": "lalu", "kelompok": {"utang_bunga": ["bank", "obligasi", "sukuk"]}}]
    pdf = angka("E", "bank", 30, 20) + angka("E", "obligasi", 10, 30) + angka("E", "total_aset", 100, 100)
    (r,) = rasio.hitung_rasio(definisi, pdf, {"E": "EPP"}, {"bank": "Utang bank"}, {"EPP": ["DES"]})
    assert r["rasio_y1"] == pytest.approx(50.0) and r["rasio_y"] is None
    assert r["hasil_rasio"] == "Y-1: 50,00% (tidak memenuhi, maks 45,00%)"
    assert "Utang bank = 20,00" in r["nilai_y1"] and r["nilai_y"] == ""
    assert rasio.hitung_rasio(definisi, pdf, {"E": "EPP"}) == []          # tanpa kriteria tambahan


def test_tambah_kriteria():
    from sikuli.rules import Rules, tambah_kriteria
    rules = Rules(rules_akun=[{"doc_type": "EPP", "kode_akun": "total_aset"}, {"doc_type": "DES", "kode_akun": "total_aset"},
                              {"doc_type": "DES", "kode_akun": "utang_bank_bunga"}])
    hasil = tambah_kriteria(rules, {"EPP": ["DES"]})
    assert [(a["doc_type"], a["kode_akun"]) for a in hasil.rules_akun if a["doc_type"] == "EPP"] == [
        ("EPP", "total_aset"), ("EPP", "utang_bank_bunga")]


def test_perhitungan_dan_sumber_angka():
    definisi = [{"no": 1, "entitas": "MI", "satuan": "persen", "rumus": "laba / aset", "teks": "ROA",
                 "kategori": "P", "field_key": "ROA", "dtype": "number", "description": ""}]
    pdf = [{"entity_key": "E", "kode_akun": "laba", "periode": "berjalan", "nilai": 5.0, "skala": 1e6,
            "halaman": 10, "label_di_dokumen": "LABA TAHUN BERJALAN", "nilai_asli": "5"},
           {"entity_key": "E", "kode_akun": "aset", "periode": "berjalan", "nilai": 100.0, "skala": 1e6,
            "halaman": 9, "label_di_dokumen": "JUMLAH ASET", "nilai_asli": "100"}]
    (r,) = rasio.hitung_rasio(definisi, pdf, {"E": "MI"}, {"laba": "Laba", "aset": "Total Aset"})
    assert r["perhitungan_y"] == "5.000.000 / 100.000.000 x 100% = 5,00%"
    assert r["sumber_y"].splitlines() == [
        '- Laba = 5.000.000,00 <- PDF hal. 10, "LABA TAHUN BERJALAN" tertulis 5 dalam jutaan Rupiah',
        '- Total Aset = 100.000.000,00 <- PDF hal. 9, "JUMLAH ASET" tertulis 100 dalam jutaan Rupiah']
    assert r["perhitungan_y1"] == "" and r["rasio_y1"] is None
