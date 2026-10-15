"""Memuat dan memvalidasi file rules (kertas kerja) di folder rules/.

Setiap file adalah CSV UTF-8 dengan baris judul. Kolom yang tidak dikenal
dibiarkan (diabaikan), kolom wajib yang hilang menghasilkan error.
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

SEMUA = {"", "*", "semua", "all"}

KOLOM_WAJIB: dict[str, list[str]] = {
    "mapping_seksi": ["kode", "seksi", "entitas", "pola"],
    "rules_akun": ["doc_type", "seksi", "kode_akun", "nama_akun"],
    "sheet_config": ["doc_type", "sheet", "seksi", "kolom_label", "kolom_nilai", "baris_mulai"],
    "mapping_akun_excel": ["doc_type", "label_excel", "kode_akun"],
    "daftar_entitas": ["entity_id", "doc_type"],
}

# File yang boleh tidak ada.
OPSIONAL = {"mapping_akun_excel", "daftar_entitas"}


def berlaku_untuk(nilai_kolom: str | None, doc_type: str | None) -> bool:
    """True bila kolom entitas/doc_type ('Semua', 'RD', 'EPP,MI') mencakup doc_type."""
    nilai = (nilai_kolom or "").strip().lower()
    if nilai in SEMUA:
        return True
    if not doc_type:
        return False
    daftar = {x.strip().lower() for x in nilai.replace(";", ",").split(",")}
    return doc_type.strip().lower() in daftar


def _aktif(baris: dict) -> bool:
    return (baris.get("aktif") or "1").strip().lower() not in {"0", "false", "tidak", "n", "no"}


def baca_csv(path: str | Path) -> list[dict]:
    with open(path, encoding="utf-8-sig", newline="") as f:
        baris = []
        for no, r in enumerate(csv.DictReader(f), start=2):
            if None in r:  # kolom lebih banyak dari judul
                raise ValueError(f"{Path(path).name} baris {no}: jumlah kolom melebihi judul; "
                                 "isi yang mengandung koma harus diapit tanda kutip")
            baris.append({(k or "").strip(): (v or "").strip() for k, v in r.items()})
    return [r for r in baris if any(r.values()) and _aktif(r)]


@dataclass
class Rules:
    mapping_seksi: list[dict] = field(default_factory=list)
    rules_akun: list[dict] = field(default_factory=list)
    sheet_config: list[dict] = field(default_factory=list)
    mapping_akun_excel: list[dict] = field(default_factory=list)
    daftar_entitas: list[dict] = field(default_factory=list)

    def akun_untuk(self, doc_type: str | None, seksi: str) -> list[dict]:
        return [
            a for a in self.rules_akun
            if a["seksi"] == seksi and berlaku_untuk(a["doc_type"], doc_type)
        ]

    def entitas(self, entity_id: str) -> dict | None:
        for e in self.daftar_entitas:
            if e["entity_id"].lower() == entity_id.lower():
                return e
        return None


# Sheet Master: kolom bagian -> kode seksi. Bagian gabungan bisa berbeda per entitas.
BAGIAN_KE_SEKSI: dict[str, str | dict[str, str]] = {
    "laporan posisi keuangan": "financial_position",
    "laporan laba rugi": "income",
    "laporan arus kas": "cashflow",
    "laporan perubahan aset bersih": "rd_netassets",
    "laporan kinerja keuangan": "rd_performance",
    "catatan atas laporan keuangan calk": "calk",
    "laporan perubahan modal epp laporan posisi keuangan mi": {"EPP": "equity", "MI": "financial_position"},
    "laporan perubahan modal": "equity",
    "akun utang berbasis bunga": "financial_position",
    "akun pendapatan non halal": "income",
}

_AWAL_ALIAS = re.compile(r"(?i)(nama akun\s*[:;]|aliases\s*[:;])")
_AKHIR_ALIAS = re.compile(r"(?i)\n\s*\n|keterangan akun|penjelasan akun|ambil nilai")


def alias_dari_deskripsi(deskripsi: str | None) -> list[str]:
    """Ambil daftar nama alias dari kolom description sheet Master.

    "Nama Akun : Kas dan Setara Kas',  'Kas dan setara kas', 'Cash and cash equivalents'.\n\n..."
    -> ["Kas dan Setara Kas", "Kas dan setara kas", "Cash and cash equivalents"]
    """
    teks = deskripsi or ""
    m = _AWAL_ALIAS.search(teks)
    if not m:
        return []
    bagian = teks[m.end():]
    akhir = _AKHIR_ALIAS.search(bagian)
    if akhir:
        bagian = bagian[:akhir.start()]
    hasil = []
    for item in bagian.replace("'", "").replace('"', "").split(","):
        item = item.split(". ")[0].strip().strip(".").strip()  # buang kalimat penjelasan setelah alias
        if item and len(item) <= 150 and item not in hasil:
            hasil.append(item)
    return hasil


def seksi_dari_bagian(bagian: str | None, entitas: str | None) -> str | None:
    from .normalisasi import normalisasi_teks

    nilai = BAGIAN_KE_SEKSI.get(normalisasi_teks(bagian))
    if isinstance(nilai, dict):
        return nilai.get((entitas or "").upper())
    return nilai


def baca_master(path: str | Path, sheet: str = "Master") -> list[dict]:
    """Sheet Master (no, entitas, bagian, akun, field_key, dtype, description) -> baris rules_akun."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        baris = list(wb[sheet].iter_rows(values_only=True))
    finally:
        wb.close()
    judul = [str(v or "").strip().lower() for v in baris[0]]
    hasil = []
    for no_baris, r in enumerate(baris[1:], start=2):
        d = {judul[i]: ("" if v is None else str(v).strip()) for i, v in enumerate(r) if i < len(judul)}
        if not d.get("field_key") and not d.get("akun"):
            continue
        alias = alias_dari_deskripsi(d.get("description"))
        hasil.append({
            "no": d.get("no", ""),
            "baris_excel": no_baris,
            "doc_type": d.get("entitas", "").upper(),
            "bagian": d.get("bagian", ""),
            "seksi": seksi_dari_bagian(d.get("bagian"), d.get("entitas")) or "",
            "kode_akun": d.get("field_key", ""),
            "nama_akun": d.get("akun", ""),
            "sinonim": "|".join(a for a in alias if a.lower() != d.get("akun", "").lower()),
            "deskripsi": d.get("description", ""),
            "dtype": d.get("dtype", ""),
        })
    return hasil


def _cari_master(folder: Path) -> Path | None:
    for path in sorted(folder.glob("*.xlsx")):
        if not path.name.startswith("~$"):
            return path
    return None


def muat_rules(folder: str | Path, folder_teknis: str | Path | None = None) -> Rules:
    """Muat rules.

    - Rules akun: file Excel di `folder` (sheet Master) -- satu-satunya rules bisnis.
      Bila tidak ada Excel, dipakai rules_akun.csv (format lama, dipakai tes).
    - Konfigurasi teknis (mapping_seksi.csv, sheet_config.csv, dan opsional
      mapping_akun_excel.csv / daftar_entitas.csv) dicari di `folder` lalu `folder_teknis`.
    """
    folder = Path(folder)
    teknis = Path(folder_teknis) if folder_teknis else folder
    data = {}
    master = _cari_master(folder)
    for nama in KOLOM_WAJIB:
        if nama == "rules_akun" and master:
            data[nama] = baca_master(master)
            continue
        path = next((p for p in (folder / f"{nama}.csv", teknis / f"{nama}.csv") if p.exists()), None)
        if path is None:
            if nama in OPSIONAL:
                data[nama] = []
                continue
            raise FileNotFoundError(f"File rules wajib tidak ditemukan: {nama}.csv (di {folder} atau {teknis})")
        data[nama] = baca_csv(path)
    return Rules(**data)


def tambah_kriteria(rules: Rules, tambahan: dict[str, list[str]] | None) -> Rules:
    """Akun jenis lain ikut berlaku untuk suatu jenis entitas, mis. {"EPP": ["DES"]}: akun DES
    disalin dengan doc_type EPP (akun yang kode_akun-nya sudah ada untuk EPP tidak diduplikasi)."""
    if not tambahan:
        return rules
    akun = list(rules.rules_akun)
    for jenis, daftar in tambahan.items():
        ada = {a["kode_akun"] for a in akun if berlaku_untuk(a["doc_type"], jenis)}
        for sumber in daftar:
            for a in rules.rules_akun:
                if a["doc_type"] == sumber and a["kode_akun"] not in ada:
                    akun.append({**a, "doc_type": jenis, "kriteria": sumber})
                    ada.add(a["kode_akun"])
    return Rules(rules.mapping_seksi, akun, rules.sheet_config, rules.mapping_akun_excel, rules.daftar_entitas)


def validasi_rules(rules: Rules) -> list[str]:
    """Kembalikan daftar masalah; list kosong berarti rules siap dipakai."""
    masalah: list[str] = []
    for nama, kolom in KOLOM_WAJIB.items():
        isi = getattr(rules, nama)
        for i, baris in enumerate(isi, start=2):
            kosong = [k for k in kolom if not baris.get(k)]
            if kosong:
                masalah.append(f"{nama}.csv baris {i}: kolom wajib kosong/tidak ada {kosong}")

    for a in rules.rules_akun:
        if a.get("bagian") is not None and not a.get("seksi"):
            masalah.append(f"Master baris {a.get('baris_excel')}: bagian '{a.get('bagian')}' untuk entitas "
                           f"'{a.get('doc_type')}' belum dikenali (tambahkan di BAGIAN_KE_SEKSI)")
    kode_seksi = {m["kode"] for m in rules.mapping_seksi}
    for nama in ("rules_akun", "sheet_config"):
        for i, baris in enumerate(getattr(rules, nama), start=2):
            if baris.get("seksi") and baris["seksi"].split(":")[0] not in kode_seksi:
                masalah.append(
                    f"{nama}.csv baris {i}: seksi '{baris['seksi']}' tidak ada di mapping_seksi"
                )

    dilihat: dict[tuple, int] = {}
    for i, a in enumerate(rules.rules_akun, start=2):
        kunci = (a.get("doc_type", "").lower(), a.get("kode_akun"))
        if kunci in dilihat:
            masalah.append(
                f"rules_akun.csv baris {i}: kode_akun '{a.get('kode_akun')}' duplikat "
                f"untuk doc_type '{a.get('doc_type')}' (lihat baris {dilihat[kunci]})"
            )
        dilihat[kunci] = i

    kode_akun = {a.get("kode_akun") for a in rules.rules_akun}
    for i, m in enumerate(rules.mapping_akun_excel, start=2):
        if m.get("kode_akun") and m["kode_akun"] not in kode_akun:
            masalah.append(
                f"mapping_akun_excel.csv baris {i}: kode_akun '{m['kode_akun']}' tidak ada di rules_akun"
            )

    for i, s in enumerate(rules.sheet_config, start=2):
        if s.get("skala") and "!" not in s["skala"] and not s["skala"].replace(".", "", 1).isdigit():
            masalah.append(f"sheet_config.csv baris {i}: skala harus angka atau rujukan sel (mis. 1000000!B31)")
        for kolom in ("baris_mulai", "baris_akhir"):
            if s.get(kolom) and not s[kolom].replace(".", "", 1).isdigit():
                masalah.append(f"sheet_config.csv baris {i}: {kolom} harus angka, isinya '{s[kolom]}'")
    return masalah
