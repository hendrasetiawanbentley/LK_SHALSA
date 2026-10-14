"""Membaca LK Excel tanpa AI: SHEET_CONFIG + mapping akun (slide 4)."""

from __future__ import annotations

import re

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from .akun import PemetaAkun
from .normalisasi import parse_angka
from .rules import Rules, berlaku_untuk


def _cocok_sheet(pola: str, nama_sheet: str) -> bool:
    pola = pola.strip()
    if pola.lower().startswith("re:"):
        return re.search(pola[3:], nama_sheet, re.IGNORECASE) is not None
    return pola.lower() == nama_sheet.strip().lower()


def _sel(baris: tuple, kolom: str | None):
    if not kolom:
        return None
    idx = column_index_from_string(kolom.strip().upper()) - 1
    return baris[idx] if idx < len(baris) else None


def baca_excel(path: str, info: dict, rules: Rules, pemeta: PemetaAkun) -> tuple[list[dict], list[str]]:
    """Baca satu file Excel. `info` hasil entitas.info_dari_nama_file.

    Kembalikan (baris_angka, peringatan). Baris yang labelnya tidak terpetakan
    tetap dikembalikan dengan kode_akun None supaya bisa ditinjau dan
    ditambahkan ke mapping_akun_excel.csv.
    """
    doc_type = info.get("doc_type")
    konfigurasi = [s for s in rules.sheet_config if berlaku_untuk(s["doc_type"], doc_type)]
    peringatan: list[str] = []
    if not konfigurasi:
        return [], [f"{info['nama_file']}: tidak ada sheet_config untuk doc_type '{doc_type}'"]

    # data_only=True: ambil nilai hasil rumus yang tersimpan, bukan rumusnya.
    wb = load_workbook(path, read_only=True, data_only=True)
    hasil: list[dict] = []
    try:
        for cfg in konfigurasi:
            sheets = [ws for ws in wb.worksheets if _cocok_sheet(cfg["sheet"], ws.title)]
            if not sheets:
                peringatan.append(f"{info['nama_file']}: sheet '{cfg['sheet']}' tidak ditemukan")
                continue
            mulai = int(cfg["baris_mulai"])
            akhir = int(cfg["baris_akhir"]) if cfg.get("baris_akhir") else None
            skala = float(cfg.get("skala") or 1)
            kolom_periode = [("berjalan", cfg["kolom_nilai"]), ("lalu", cfg.get("kolom_nilai_lalu"))]

            for ws in sheets:
                for no_baris, baris in enumerate(
                    ws.iter_rows(min_row=mulai, max_row=akhir, values_only=True), start=mulai
                ):
                    label = _sel(baris, cfg["kolom_label"])
                    if label is None or str(label).strip() == "":
                        continue
                    label = str(label).strip()
                    nilai_periode = []
                    for periode, kolom in kolom_periode:
                        mentah = _sel(baris, kolom)
                        nilai = parse_angka(mentah)
                        if nilai is not None:
                            nilai_periode.append((periode, nilai * skala, mentah))
                    if not nilai_periode:
                        continue  # baris judul / subjudul tanpa angka
                    kode, cara = pemeta.petakan(label, doc_type, cfg["seksi"], ws.title)
                    for periode, nilai, mentah in nilai_periode:
                        hasil.append({
                            "sumber_file": info["nama_file"],
                            "entity_key": info["entity_key"],
                            "entity_id": info["entity_id"],
                            "doc_type": doc_type,
                            "tahun": info.get("tahun"),
                            "sheet": ws.title,
                            "baris": no_baris,
                            "seksi": cfg["seksi"],
                            "label_asli": label,
                            "kode_akun": kode,
                            "cara_pemetaan": cara,
                            "periode": periode,
                            "nilai": nilai,
                            "nilai_asli": None if mentah is None else str(mentah),
                            "skala": skala,
                        })
    finally:
        wb.close()
    return hasil, peringatan
