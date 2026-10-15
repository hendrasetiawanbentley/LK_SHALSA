"""Pipeline lengkap: PDF + Excel -> hasil komparasi (.xlsx, .csv, ringkasan.md).

    PYTHONPATH=src python -m sikuli.lokal            # baca input/pdf dan input/excel, tulis ke hasil/

Dijalankan oleh GitHub Actions (.github/workflows/analisis.yml).
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from . import ai, ocr, rasio
from .akun import PemetaAkun
from .config import muat_config
from .ekstraksi_lokal import ekstraksi_seksi
from .entitas import info_dari_nama_file
from .excel import baca_excel
from .komparasi import MATCH, bandingkan
from .pdf_lokal import halaman_dari_pdf
from .pipeline import buat_permintaan_ekstraksi, konteks_permintaan, proses_halaman
from .rules import Rules, berlaku_untuk, muat_rules, tambah_kriteria, validasi_rules
from .seksi import kompilasi_pola

ROOT = Path(__file__).resolve().parents[2]


def format_angka(x: float | None) -> str:
    """1250 -> '1,250'; 1250.4 -> '1,250.40'; None -> '-'."""
    if x is None:
        return "-"
    return f"{x:,.0f}" if float(x).is_integer() else f"{x:,.2f}"


def catatan_otomatis(baris: dict, catatan_match: str = "Nilai sesuai") -> str:
    """Catatan deterministik per baris (pengganti ai_gen)."""
    status = baris["status"]
    if status == MATCH:
        if baris.get("indikasi") == "BEDA_TANDA":
            return f"{catatan_match} (hanya beda tanda penyajian; sesuai rules Master bukan perbedaan nilai)"
        return catatan_match
    if status == "PDF_ONLY":
        return "Akun ada di PDF tetapi tidak ditemukan di Excel; periksa sheet_config atau mapping label Excel."
    if status == "EXCEL_ONLY":
        return ("Akun ada di Excel tetapi tidak terbaca di PDF; periksa apakah akun disajikan di laporan "
                "atau tambahkan sinonim label PDF di rules_akun.")
    indikasi = baris.get("indikasi")
    if indikasi == "BEDA_TANDA" and status == "DIFFERENT":
        persen = baris.get("selisih_persen")
        return (f"Tanda berlawanan dan nilai absolut juga berbeda {format_angka(baris['selisih'])} "
                f"{baris.get('satuan_pdf') or ''}" + (f" ({persen:.2f}%)" if persen is not None else "")
                + "; perlu konfirmasi ke entitas.").replace("  ", " ")
    if indikasi == "BEDA_SKALA":
        return "Nilai berbeda kelipatan ribuan/jutaan; periksa satuan penyajian PDF atau skala di sheet_config."
    persen = baris.get("selisih_persen")
    persen_teks = f" ({persen:.2f}%)" if persen is not None else ""
    return (f"Selisih {format_angka(baris['selisih'])} {baris.get('satuan_pdf') or ''}{persen_teks}; "
            "perlu konfirmasi ke entitas.").replace("  ", " ")


def _ekstraksi_aturan(teks: list[dict], cfg: dict, rules: Rules, peringatan: list[str]) -> list[dict]:
    """Ekstraksi akun tanpa AI dari baris teks_pdf satu dokumen."""
    per_seksi: dict[tuple, dict] = {}
    for p in buat_permintaan_ekstraksi(teks, cfg, rules):
        kunci = (p["doc_id"], p["seksi"])
        if kunci in per_seksi:
            per_seksi[kunci]["akun"].extend(p["akun"])
        else:
            per_seksi[kunci] = {**p, "akun": list(p["akun"])}
    angka = []
    for (doc_id, seksi), p in per_seksi.items():
        halaman = [t for t in teks if t["doc_id"] == doc_id and t["seksi"] == seksi]
        baris, tidak_ada = ekstraksi_seksi(halaman, p["akun"], konteks_permintaan(p), p["zona_atas"])
        for b in baris:
            b["sumber_ekstraksi"], b["verifikasi"] = "aturan", "cocok"
        angka.extend(baris)
        if tidak_ada:
            peringatan.append(f"{p['sumber_file']} / {seksi}: akun tidak ditemukan di PDF: {', '.join(tidak_ada)}")
    return angka


AKHIRAN_PERIODE_LALU = ("_awal_periode", "_periode_sebelumnya")


def turunan_periode_lalu(baris: list[dict], rules: Rules, doc_type: str | None) -> int:
    """Akun `<akun>_awal_periode` / `<akun>_periode_sebelumnya` (mis. NAB awal periode) diambil
    dari kolom periode lalu akun dasarnya, bila Excel tidak punya baris sendiri untuk akun itu."""
    ada = {b["kode_akun"] for b in baris if b.get("kode_akun")}
    tambahan = []
    for a in rules.rules_akun:
        kode = a["kode_akun"]
        akhiran = next((x for x in AKHIRAN_PERIODE_LALU if kode.endswith(x)), None)
        if not akhiran or kode in ada or not berlaku_untuk(a["doc_type"], doc_type):
            continue
        dasar = kode[: -len(akhiran)]
        for b in baris:
            if b.get("kode_akun") == dasar and b["periode"] == "lalu":
                tambahan.append({**b, "kode_akun": kode, "periode": "berjalan",
                                 "cara_pemetaan": "turunan_periode_lalu",
                                 "label_asli": f"{b['label_asli']} (kolom periode lalu)"})
                break
    baris.extend(tambahan)
    return len(tambahan)


def _baca_satu(path: str, zona: int):
    try:
        return halaman_dari_pdf(path, zona)
    except Exception as e:  # dikembalikan, bukan dilempar, agar PDF lain tetap diproses
        return e


def _baca_pdf_paralel(pdf_paths: list[str], zona: int) -> dict:
    """Teks semua PDF dibaca paralel (pdfplumber lambat untuk LK ratusan halaman)."""
    if len(pdf_paths) < 2:
        return {p: _baca_satu(p, zona) for p in pdf_paths}
    from concurrent.futures import ProcessPoolExecutor
    with ProcessPoolExecutor(max_workers=min(len(pdf_paths), os.cpu_count() or 1)) as ex:
        return dict(zip(pdf_paths, ex.map(_baca_satu, pdf_paths, [zona] * len(pdf_paths))))


def _jalankan_ocr(klien, cfg_ai: dict, path: str, halaman, zona: int, peringatan: list[str]) -> None:
    """Halaman bergambar tanpa teks (scan) dibaca Claude bila ai.ocr aktif; selain itu hanya dicatat."""
    nomor = ocr.perlu_ocr(halaman)
    if not nomor:
        return
    nama, daftar = Path(path).name, ", ".join(map(str, nomor))
    if not (klien and cfg_ai.get("ocr")):
        peringatan.append(f"{nama}: {len(nomor)} halaman bergambar tanpa teks (hal. {daftar}); "
                          "aktifkan ai.ocr untuk OCR dengan Claude")
        return
    try:
        hasil = ocr.ocr_halaman(klien, path, nomor, int(cfg_ai.get("ocr_halaman_per_permintaan", 10)))
    except Exception as e:  # API error / rekaman OCR tidak ada
        peringatan.append(f"{nama}: OCR gagal ({type(e).__name__}: {e}); hal. {daftar} tanpa teks")
        return
    n = ocr.terapkan_ocr(halaman, hasil, zona)
    peringatan.append(f"{nama}: {n} dari {len(nomor)} halaman scan dibaca dengan OCR Claude (sumber_teks = ocr_claude)")


def _proses_pdf_ai(klien, cfg: dict, rules: Rules, path: str, info: dict, halaman,
                   lokasi: list, angka_pdf: list, peringatan: list) -> bool:
    """Ekstraksi PDF oleh Claude. True = dokumen selesai (berhasil, atau gagal dengan ai.wajib)."""
    cfg_ai = cfg.get("ai") or {}
    if not (klien and cfg_ai.get("ekstraksi_pdf", True) and info["entitas_dikenal"]):
        return False
    nama = Path(path).name
    try:
        lok, ang, pesan = ai.ekstraksi_pdf(klien, path, info, rules, cfg, [h.teks for h in halaman])
    except Exception as e:  # API error, refusal, PDF terlalu besar
        if cfg_ai.get("wajib"):
            peringatan.append(f"{nama}: GAGAL - ekstraksi AI error ({type(e).__name__}: {e}); "
                              "dokumen tidak diproses (ai.wajib = true)")
            return True
        peringatan.append(f"{nama}: ekstraksi AI gagal ({type(e).__name__}: {e}); memakai aturan")
        return False
    lokasi.extend(lok)
    angka_pdf.extend(ang)
    peringatan.extend(pesan)
    return True


def jalankan(pdf_paths: list[str], excel_paths: list[str], cfg: dict, rules: Rules,
             klien: ai.KlienAI | None = None, definisi_rasio: list[dict] | None = None) -> dict:
    """Pipeline lengkap. `klien` = KlienAI untuk memakai Claude; None = hanya aturan."""
    cfg_ai = cfg.get("ai") or {}
    pola = kompilasi_pola(rules.mapping_seksi)
    teks, lokasi, angka_pdf, peringatan, teks_halaman = [], [], [], [], []
    zona = int(cfg["deteksi_seksi"]["baris_zona_atas"])
    semua_halaman = _baca_pdf_paralel(pdf_paths, zona)
    for path in pdf_paths:
        nama = Path(path).name
        halaman = semua_halaman[path]
        if isinstance(halaman, Exception):  # PDF rusak / terkunci
            peringatan.append(f"{nama}: gagal dibaca ({type(halaman).__name__}: {halaman})")
            continue
        info = info_dari_nama_file(path, cfg, rules)
        _jalankan_ocr(klien, cfg_ai, path, halaman, zona, peringatan)
        if not _proses_pdf_ai(klien, cfg, rules, path, info, halaman, lokasi, angka_pdf, peringatan):
            hasil = proses_halaman(path, halaman, cfg, rules, pola)
            teks.extend(hasil["teks"])
            lokasi.extend(hasil["lokasi"])
            peringatan.extend(hasil["peringatan"])
            angka_pdf.extend(_ekstraksi_aturan(hasil["teks"], cfg, rules, peringatan))
        teks_halaman.extend(ocr.baris_teks(info, halaman, [l for l in lokasi if l.get("nama_file") == info["nama_file"]]))

    angka_excel, pemeta = [], PemetaAkun(rules)
    for path in excel_paths:
        info = info_dari_nama_file(path, cfg, rules)
        if not info["entitas_dikenal"]:
            peringatan.append(f"{info['nama_file']}: entity_id/doc_type tidak dikenali dari nama file, dilewati")
            continue
        try:
            baris, pesan = baca_excel(path, info, rules, pemeta)
        except Exception as e:
            peringatan.append(f"{info['nama_file']}: gagal dibaca ({type(e).__name__}: {e})")
            continue
        if klien and cfg_ai.get("pemetaan_excel", True):
            try:
                n = ai.pemetaan_excel(klien, baris, rules, info["doc_type"])
                if n:
                    pesan.append(f"{info['nama_file']}: {n} label dipetakan oleh AI (cara_pemetaan = ai), mohon ditinjau")
            except Exception as e:
                pesan.append(f"{info['nama_file']}: pemetaan AI gagal ({type(e).__name__}: {e})")
        turunan_periode_lalu(baris, rules, info["doc_type"])
        angka_excel.extend(baris)
        peringatan.extend(pesan)

    nama_akun = {a["kode_akun"]: a["nama_akun"] for a in rules.rules_akun}
    komparasi = bandingkan(angka_pdf, angka_excel, cfg["komparasi"], nama_akun)
    urutan = {"DIFFERENT": 0, "PDF_ONLY": 1, "EXCEL_ONLY": 2, MATCH: 3}
    komparasi.sort(key=lambda b: (urutan[b["status"]], b["entity_key"], b["kode_akun"]))
    for b in komparasi:
        b["catatan"] = catatan_otomatis(b, cfg["komparasi"].get("catatan_match", "Nilai sesuai"))
        b["sumber_catatan"] = "otomatis"

    kesimpulan = []
    if klien and cfg_ai.get("catatan_dan_kesimpulan", True):
        for entity_key in sorted({b["entity_key"] for b in komparasi}):
            try:
                kesimpulan.append(ai.analisis_entitas(
                    klien, entity_key, [b for b in komparasi if b["entity_key"] == entity_key]))
            except Exception as e:
                peringatan.append(f"{entity_key}: analisis AI gagal ({type(e).__name__}: {e})")
    hasil_rasio = []
    if definisi_rasio:
        # Rasio hanya dari angka PDF (laporan audited).
        entitas = {b["entity_key"]: b["doc_type"] for b in angka_pdf if b.get("doc_type")}
        hasil_rasio = rasio.hitung_rasio(definisi_rasio, angka_pdf, entitas, nama_akun,
                                         (cfg.get("entitas") or {}).get("kriteria_tambahan"))
        if klien and cfg_ai.get("analisis_rasio", True):
            for entity_key in sorted(entitas):
                baris = [r for r in hasil_rasio if r["entity_key"] == entity_key]
                if not baris:
                    continue
                try:
                    ai.analisis_rasio(klien, entity_key, baris)
                except Exception as e:
                    peringatan.append(f"{entity_key}: analisis rasio AI gagal ({type(e).__name__}: {e})")
    return {"komparasi": komparasi, "angka_pdf": angka_pdf, "angka_excel": angka_excel, "rasio": hasil_rasio,
            "lokasi_seksi": lokasi, "teks_pdf": teks, "teks_halaman": teks_halaman, "peringatan": peringatan,
            "ringkasan": ringkasan(komparasi), "kesimpulan": kesimpulan,
            "mode": f"AI ({klien.model})" if klien else "aturan (tanpa AI)"}


def ringkasan(komparasi: list[dict]) -> list[dict]:
    per: dict[str, dict] = {}
    for b in komparasi:
        r = per.setdefault(b["entity_key"], {"entity_key": b["entity_key"], "MATCH": 0, "DIFFERENT": 0,
                                             "PDF_ONLY": 0, "EXCEL_ONLY": 0})
        r[b["status"]] += 1
    for r in per.values():
        total = r["MATCH"] + r["DIFFERENT"] + r["PDF_ONLY"] + r["EXCEL_ONLY"]
        r["persen_match"] = round(100 * r["MATCH"] / total, 1) if total else 0.0
    return sorted(per.values(), key=lambda r: r["persen_match"])


KOLOM = {
    "Ringkasan": ["entity_key", "MATCH", "DIFFERENT", "PDF_ONLY", "EXCEL_ONLY", "persen_match"],
    "Komparasi": ["entity_key", "seksi", "kode_akun", "nama_akun", "periode", "satuan_pdf", "nilai_pdf",
                  "nilai_excel", "nilai_excel_asli", "selisih", "selisih_persen", "status", "indikasi", "catatan", "sumber_catatan", "halaman_pdf",
                  "verifikasi_pdf", "sumber_pdf",
                  "sumber_excel", "label_pdf", "label_excel"],
    "Angka_PDF": ["sumber_file", "entity_key", "seksi", "kode_akun", "label_di_dokumen", "periode", "nilai",
                  "nilai_asli", "skala", "halaman", "sumber_ekstraksi", "verifikasi"],
    "Angka_Excel": ["sumber_file", "entity_key", "sheet", "baris", "seksi", "label_asli", "kode_akun",
                    "cara_pemetaan", "periode", "nilai", "nilai_asli"],
    "Lokasi_Seksi": ["nama_file", "entity_key", "seksi", "halaman_mulai", "halaman_akhir", "judul_cocok"],
    "Kesimpulan": ["entity_key", "kesimpulan", "rekomendasi"],
    "Rasio": rasio.KOLOM,
    "Teks_PDF": ["entity_key", "nama_file", "halaman", "no_baris", "seksi", "sumber_teks", "teks"],
}
WARNA = {"MATCH": "D9F2E1", "DIFFERENT": "FBDAD7", "PDF_ONLY": "FFF1C2", "EXCEL_ONLY": "FFF1C2"}


def tulis_excel(hasil: dict, path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    sumber = {"Ringkasan": "ringkasan", "Komparasi": "komparasi", "Angka_PDF": "angka_pdf",
              "Angka_Excel": "angka_excel", "Lokasi_Seksi": "lokasi_seksi", "Kesimpulan": "kesimpulan",
              "Rasio": "rasio", "Teks_PDF": "teks_halaman"}
    for sheet, kolom in KOLOM.items():
        ws = wb.create_sheet(sheet)
        ws.append(kolom)
        for c in ws[1]:
            c.font = Font(bold=True)
        for b in hasil.get(sumber[sheet], []):
            ws.append([b.get(k) for k in kolom])
            if sheet == "Komparasi" and b["status"] in WARNA:
                ws.cell(ws.max_row, kolom.index("status") + 1).fill = PatternFill("solid", fgColor=WARNA[b["status"]])
        ws.freeze_panes = "A2"
    ws = wb.create_sheet("Peringatan")
    ws.append(["peringatan"])
    for p in hasil["peringatan"]:
        ws.append([p])
    wb.save(path)


def _kumpulkan(folder: str | None, akhiran: tuple[str, ...]) -> list[str]:
    if not folder:
        return []
    return sorted(str(p) for p in Path(folder).rglob("*")
                  if p.suffix.lower() in akhiran and not p.name.startswith("~$"))


def tulis_csv(hasil: dict, folder: Path) -> None:
    sumber = {"ringkasan": "Ringkasan", "komparasi": "Komparasi", "angka_pdf": "Angka_PDF",
              "angka_excel": "Angka_Excel", "lokasi_seksi": "Lokasi_Seksi", "kesimpulan": "Kesimpulan",
              "rasio": "Rasio", "teks_pdf": "Teks_PDF"}
    for kunci, sheet in sumber.items():
        with open(folder / f"{kunci}.csv", "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(KOLOM[sheet])
            data = hasil.get("teks_halaman" if kunci == "teks_pdf" else kunci, [])
            w.writerows([b.get(k) for k in KOLOM[sheet]] for b in data)
    with open(folder / "peringatan.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["peringatan"])
        w.writerows([p] for p in hasil["peringatan"])


def tulis_json(hasil: dict, path) -> None:
    """Semua hasil (tanpa teks PDF mentah) dalam satu file JSON, dipakai scripts/buat_dashboard.py."""
    data = {k: v for k, v in hasil.items() if k not in ("teks_pdf", "teks_halaman")}
    Path(path).write_text(json.dumps(data, ensure_ascii=False, default=str, indent=1), encoding="utf-8")


def ringkasan_markdown(hasil: dict, jumlah_pdf: int, jumlah_excel: int) -> str:
    """Ringkasan untuk halaman run GitHub Actions (GITHUB_STEP_SUMMARY)."""
    angka = format_angka

    def sel(x):
        return "" if x is None else str(x).replace("|", "\\|")

    k = hasil["komparasi"]
    status = {s: sum(1 for b in k if b["status"] == s) for s in ("MATCH", "DIFFERENT", "PDF_ONLY", "EXCEL_ONLY")}
    md = [f"## Hasil komparasi SIKULI", "",
          f"Mode: **{hasil.get('mode', 'aturan (tanpa AI)')}**. "
          f"{jumlah_pdf} PDF, {jumlah_excel} Excel, {len(k)} baris akun. "
          + ", ".join(f"**{v}** {s}" for s, v in status.items()), "",
          "### Ringkasan per entitas", "",
          "| entity_key | MATCH | DIFFERENT | PDF_ONLY | EXCEL_ONLY | % match |", "|---|--:|--:|--:|--:|--:|"]
    md += [f"| {sel(r['entity_key'])} | {r['MATCH']} | {r['DIFFERENT']} | {r['PDF_ONLY']} | {r['EXCEL_ONLY']} | "
           f"{r['persen_match']} |" for r in hasil["ringkasan"]]
    if hasil.get("kesimpulan"):
        md += ["", "### Kesimpulan dan rekomendasi", ""]
        for r in hasil["kesimpulan"]:
            md += [f"**{sel(r['entity_key'])}**", "", r["kesimpulan"], ""]
            if r["rekomendasi"]:
                md += [r["rekomendasi"], ""]
    perlu = [b for b in k if b["status"] != MATCH]
    md += ["", f"### Perlu ditinjau ({len(perlu)} baris)", ""]
    if perlu:
        md += ["Nilai dalam satuan PDF; nilai Excel sudah dibagi sesuai satuan tersebut.", "",
               "| entitas | akun | periode | satuan | nilai PDF | nilai Excel | selisih | status | catatan |",
               "|---|---|---|---|--:|--:|--:|---|---|"]
        md += [f"| {sel(b['entity_key'])} | {sel(b['kode_akun'])} | {b['periode']} | {b['satuan_pdf']} | "
               f"{angka(b['nilai_pdf'])} | "
               f"{angka(b['nilai_excel'])} | {angka(b['selisih'])} | {b['status']} | {sel(b['catatan'])} |"
               for b in perlu[:200]]
        if len(perlu) > 200:
            md.append(f"\n_{len(perlu) - 200} baris lain ada di file hasil._")
    else:
        md.append("Semua akun MATCH.")
    if hasil["peringatan"]:
        md += ["", f"### Peringatan ({len(hasil['peringatan'])})", ""] + [f"- {p}" for p in hasil["peringatan"]]
    md += ["", "File lengkap (`hasil_komparasi.xlsx` dan CSV) ada di **Artifacts** di bawah halaman run ini."]
    return "\n".join(md) + "\n"


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="SIKULI: komparasi LK PDF vs Excel")
    ap.add_argument("--pdf", default=str(ROOT / "input" / "pdf"), help="folder berisi LK PDF")
    ap.add_argument("--excel", default=str(ROOT / "input" / "excel"), help="folder berisi LK Excel")
    ap.add_argument("--rules", default=str(ROOT / "rules"))
    ap.add_argument("--config", default=str(ROOT / "config" / "pipeline.yaml"))
    ap.add_argument("--out-dir", default=str(ROOT / "hasil"), help="folder hasil (xlsx, csv, md)")
    ap.add_argument("--ai-rekaman", default=None,
                    help="folder jawaban AI yang sudah direkam (lihat ai.KlienRekaman); tanpa API key")
    args = ap.parse_args(argv)

    rules = muat_rules(args.rules, Path(args.config).parent)
    masalah = validasi_rules(rules)
    if masalah:
        print("Rules tidak valid:\n" + "\n".join(f"- {m}" for m in masalah))
        return 1
    pdf = _kumpulkan(args.pdf, (".pdf",))
    excel = _kumpulkan(args.excel, (".xlsx", ".xlsm"))
    out = Path(args.out_dir)
    out.mkdir(parents=True, exist_ok=True)
    if not pdf and not excel:
        pesan = (f"## Hasil komparasi SIKULI\n\nBelum ada file: unggah PDF ke `input/pdf/` dan "
                 f"Excel ke `input/excel/`.\n")
        (out / "ringkasan.md").write_text(pesan, encoding="utf-8")
        print(pesan)
        return 0

    cfg = muat_config(args.config)
    rules = tambah_kriteria(rules, (cfg.get("entitas") or {}).get("kriteria_tambahan"))
    cfg_ai = cfg.get("ai") or {}
    if args.ai_rekaman:
        klien = ai.KlienRekaman(args.ai_rekaman, model=f"rekaman {args.ai_rekaman}")
    else:
        klien = ai.buat_klien(cfg)
    if cfg_ai.get("aktif") and klien is None:
        if cfg_ai.get("wajib"):
            pesan = ("## Hasil komparasi SIKULI\n\n**Gagal: AI wajib dipakai tetapi secret `ANTHROPIC_API_KEY` "
                     "belum diisi.**\n\nIsi di GitHub: Settings → Secrets and variables → Actions → "
                     "New repository secret, nama `ANTHROPIC_API_KEY`. Lalu jalankan ulang workflow "
                     "(Actions → analisis → Run workflow).\n")
            (out / "ringkasan.md").write_text(pesan, encoding="utf-8")
            print(pesan)
            return 1
        print("AI aktif di config tetapi ANTHROPIC_API_KEY tidak ada; memakai aturan saja.")
    folder_config = Path(args.config).parent
    definisi_rasio = rasio.muat_rasio(folder_config / "rasio_kk_template.xlsx", folder_config / "rasio.yaml", args.rules)
    hasil = jalankan(pdf, excel, cfg, rules, klien, definisi_rasio)
    tulis_excel(hasil, out / "hasil_komparasi.xlsx")
    if hasil["rasio"]:
        rasio.tulis_kk(hasil["rasio"], out / "Rasio_KK.xlsx")
    tulis_csv(hasil, out)
    tulis_json(hasil, out / "hasil.json")
    md = ringkasan_markdown(hasil, len(pdf), len(excel))
    (out / "ringkasan.md").write_text(md, encoding="utf-8")
    print(md)
    print(f"Hasil ditulis ke {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
