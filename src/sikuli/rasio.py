"""Rasio kertas kerja (Sheet3 rules) dari angka PDF (diaudit) dan Excel.

Definisi rasio = Sheet3 (No, Kategori, Entitas, field_key, dtype, Description) digabung dengan
rumus di config/rasio.yaml. Perhitungan deterministik; analisis per rasio ditulis AI
(ai.analisis_rasio) dan masuk kolom "Analisis AI".
"""

from __future__ import annotations

import ast
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .komparasi import _indeks

KOLOM_KK = ["No", "Kategori", "Entitas", "field_key", "dtype", "Description", "Nilai Y-1", "Nilai Y",
            "Rumus Rasio AI", "Analisis AI"]
KOLOM_TAMBAHAN = ["Satuan", "Nilai Y (Excel)", "Komponen Y (PDF)", "Status ambang", "Keterangan"]
KOLOM = ["entity_key", "no", "kategori", "entitas", "field_key", "dtype", "description", "nilai_y1", "nilai_y",
         "rumus_ai", "analisis_ai", "sumber_analisis", "satuan", "nilai_y_excel", "komponen_y", "status_ambang",
         "keterangan"]
FUNGSI = {"rata2", "opsional", "abs", "min", "max"}


class AkunHilang(Exception):
    pass


def _baca_sheet3(path: Path) -> dict[tuple[int, str], dict]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Sheet3" not in wb.sheetnames:
            return {}
        baris = list(wb["Sheet3"].iter_rows(values_only=True))
    finally:
        wb.close()
    kepala = [str(k).strip() if k else "" for k in baris[0]]
    hasil = {}
    for r in baris[1:]:
        d = dict(zip(kepala, r))
        if d.get("No") is None or not d.get("Entitas"):
            continue
        kunci = (int(d["No"]), str(d["Entitas"]).strip())
        hasil[kunci] = {"kategori": str(d.get("Kategori") or "").strip(), "field_key": str(d.get("field_key") or "").strip(),
                        "dtype": d.get("dtype") or "number", "description": d.get("Description") or ""}
    return hasil


def muat_rasio(folder_rules: str | Path, path_yaml: str | Path) -> list[dict]:
    """Definisi rasio: rumus dari yaml + kategori/nama/deskripsi dari Sheet3 file rules."""
    path_yaml = Path(path_yaml)
    if not path_yaml.exists():
        return []
    sheet3 = {}
    master = next(iter(sorted(Path(folder_rules).glob("*.xlsx"))), None)
    if master:
        sheet3 = _baca_sheet3(master)
    hasil = []
    for r in yaml.safe_load(path_yaml.read_text(encoding="utf-8")).get("rasio", []):
        _cek_rumus(r["rumus"])
        info = sheet3.get((int(r["nomor"]), r["entitas"]), {})
        hasil.append({**r, "no": int(r["nomor"]), "kategori": info.get("kategori", r.get("kategori", "")),
                      "field_key": info.get("field_key", r.get("nama", f"Rasio {r['nomor']}")),
                      "dtype": info.get("dtype", "number"), "description": info.get("description", "")})
    return hasil


def _cek_rumus(rumus: str) -> None:
    """Rumus hanya boleh berisi angka, nama akun, operator aritmetika, dan fungsi yang diizinkan."""
    izin = (ast.Expression, ast.BinOp, ast.UnaryOp, ast.Name, ast.Load, ast.Constant, ast.Call,
            ast.Add, ast.Sub, ast.Mult, ast.Div, ast.USub, ast.UAdd)
    for n in ast.walk(ast.parse(rumus, mode="eval")):
        if not isinstance(n, izin):
            raise ValueError(f"rumus rasio tidak diizinkan: {rumus!r} ({type(n).__name__})")
        if isinstance(n, ast.Call) and not (isinstance(n.func, ast.Name) and n.func.id in FUNGSI):
            raise ValueError(f"fungsi tidak diizinkan di rumus: {rumus!r}")


def _nama_akun(rumus: str) -> list[str]:
    nama = []
    for n in ast.walk(ast.parse(rumus, mode="eval")):
        if isinstance(n, ast.Name) and n.id not in FUNGSI and n.id not in nama:
            nama.append(n.id)
    return nama


def evaluasi(rumus: str, nilai) -> float:
    """`nilai(nama)` -> float atau None. Akun yang hilang -> AkunHilang (kecuali dalam opsional())."""
    def ev(n):
        if isinstance(n, ast.Expression):
            return ev(n.body)
        if isinstance(n, ast.Constant):
            return float(n.value)
        if isinstance(n, ast.Name):
            v = nilai(n.id)
            if v is None:
                raise AkunHilang(n.id)
            return v
        if isinstance(n, ast.UnaryOp):
            return -ev(n.operand) if isinstance(n.op, ast.USub) else ev(n.operand)
        if isinstance(n, ast.BinOp):
            a, b = ev(n.left), ev(n.right)
            if isinstance(n.op, ast.Add):
                return a + b
            if isinstance(n.op, ast.Sub):
                return a - b
            if isinstance(n.op, ast.Mult):
                return a * b
            if b == 0:
                raise ZeroDivisionError
            return a / b
        f = n.func.id
        if f == "opsional":
            try:
                return ev(n.args[0])
            except AkunHilang:
                return 0.0
        args = [ev(a) for a in n.args]
        if f == "rata2":
            return sum(args) / len(args)
        return {"abs": abs, "min": min, "max": max}[f](*args)
    return ev(ast.parse(rumus, mode="eval"))


def _nilai_entitas(baris: list[dict], entity_key: str) -> dict[tuple[str, str], float]:
    """(kode_akun, periode) -> nilai mata uang penuh."""
    hasil = {}
    for (ek, kode, periode), b in _indeks(baris, {"berjalan", "lalu"}, "x").items():
        if ek == entity_key:
            hasil[(kode, periode)] = b["nilai"] * (b.get("skala") or 1.0)
    return hasil


def _pengambil(data: dict, periode: str):
    def ambil(nama: str):
        if nama.endswith("__lalu"):
            return data.get((nama[:-6], "lalu")) if periode == "berjalan" else None
        return data.get((nama, periode))
    return ambil


def _hitung(rumus: str, satuan: str, ambil) -> tuple[float | None, str | None]:
    try:
        v = evaluasi(rumus, ambil)
    except AkunHilang as e:
        return None, f"akun {e.args[0]} tidak tersedia"
    except ZeroDivisionError:
        return None, "pembagi bernilai nol"
    return (v * 100 if satuan == "persen" else v), None


def _fmt(x: float) -> str:
    return f"{x:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")


def hitung_rasio(definisi: list[dict], angka_pdf: list[dict], angka_excel: list[dict],
                 entitas: dict[str, str]) -> list[dict]:
    """`entitas`: entity_key -> doc_type. Satu baris per entitas x rasio yang berlaku."""
    hasil = []
    for entity_key, doc_type in sorted(entitas.items()):
        pdf = _nilai_entitas(angka_pdf, entity_key)
        excel = _nilai_entitas([{**b, "skala": 1.0} for b in angka_excel], entity_key)
        for d in (x for x in definisi if x["entitas"] == doc_type):
            y, ket_y = _hitung(d["rumus"], d["satuan"], _pengambil(pdf, "berjalan"))
            y1, ket_y1 = _hitung(d["rumus"], d["satuan"], _pengambil(pdf, "lalu"))
            ye, _ = _hitung(d["rumus"], d["satuan"], _pengambil(excel, "berjalan"))
            ambil = _pengambil(pdf, "berjalan")
            komponen = "; ".join(f"{n} = {_fmt(ambil(n))}" if ambil(n) is not None else f"{n} = tidak ada"
                                 for n in _nama_akun(d["rumus"]))
            status = None
            if d.get("ambang") is not None and y is not None:
                status = "Memenuhi" if y >= float(d["ambang"]) else "Tidak memenuhi"
            ket = []
            if ket_y:
                ket.append(f"Y: {ket_y}")
            if ket_y1:
                ket.append(f"Y-1: {ket_y1}")
            hasil.append({
                "entity_key": entity_key, "no": d["no"], "kategori": d["kategori"], "entitas": doc_type,
                "field_key": d["field_key"], "dtype": d["dtype"], "description": d["description"],
                "nilai_y1": y1, "nilai_y": y, "rumus_ai": d["teks"], "analisis_ai": "", "sumber_analisis": None,
                "satuan": {"persen": "%", "kali": "x", "nominal": "nominal"}[d["satuan"]],
                "nilai_y_excel": ye, "komponen_y": komponen, "status_ambang": status, "keterangan": "; ".join(ket),
            })
    return hasil


def nama_sheet(entity_keys: list[str]) -> dict[str, str]:
    """Nama sheet pendek dan unik: JENIS_KODE, ditambah bagian berikutnya bila bentrok."""
    hasil = {}
    for e in entity_keys:
        bagian = e.split("_")
        n = 2
        while n < len(bagian) and sum(1 for x in entity_keys if x.split("_")[:n] == bagian[:n]) > 1:
            n += 1
        hasil[e] = "_".join(bagian[:n])[:31]
    return hasil


def tulis_kk(rasio: list[dict], path) -> None:
    """Kertas kerja rasio: satu sheet per entitas dengan kolom seperti Rasio_KK, plus sheet Semua."""
    wb = Workbook()
    wb.remove(wb.active)
    kolom = KOLOM_KK + KOLOM_TAMBAHAN
    kunci = ["no", "kategori", "entitas", "field_key", "dtype", "description", "nilai_y1", "nilai_y", "rumus_ai",
             "analisis_ai", "satuan", "nilai_y_excel", "komponen_y", "status_ambang", "keterangan"]
    lebar = {"A": 5, "B": 16, "C": 8, "D": 30, "E": 8, "F": 50, "G": 16, "H": 16, "I": 40, "J": 70,
             "K": 9, "L": 16, "M": 50, "N": 14, "O": 30}
    kepala = PatternFill("solid", fgColor="DCE3F2")

    def sheet(nama: str, isi: list[dict], dengan_entitas: bool):
        ws = wb.create_sheet(nama[:31])
        ws.append((["entity_key"] if dengan_entitas else []) + kolom)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True), kepala
        for r in isi:
            ws.append(([r["entity_key"]] if dengan_entitas else []) + [r.get(k) for k in kunci])
            for c in ws[ws.max_row]:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            geser = 1 if dengan_entitas else 0
            fmt = "#,##0" if r["satuan"] == "nominal" else "#,##0.00"
            for kol in (7, 8, 12):
                ws.cell(ws.max_row, kol + geser).number_format = fmt
        for huruf, w in lebar.items():
            ws.column_dimensions[chr(ord(huruf) + (1 if dengan_entitas else 0))].width = w
        if dengan_entitas:
            ws.column_dimensions["A"].width = 40
        ws.freeze_panes = "A2"

    for e, nama in nama_sheet(sorted({r["entity_key"] for r in rasio})).items():
        sheet(nama, [r for r in rasio if r["entity_key"] == e], False)
    sheet("Semua", rasio, True)
    wb.save(path)
